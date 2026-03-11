from flask import Flask, render_template, request, redirect, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from flask_socketio import SocketIO
from PyPDF2 import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
import io
import datetime
import os
from openpyxl import Workbook
from sqlalchemy import or_   # ✅ Added for multi-field search
from werkzeug.middleware.proxy_fix import ProxyFix

app = Flask(__name__)
app.wsgi_app = ProxyFix(app.wsgi_app, x_proto=1, x_host=1)

app.secret_key = os.environ.get("SECRET_KEY", "supersecretkey")

app.config["SESSION_COOKIE_SECURE"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "None"
app.config["SESSION_COOKIE_HTTPONLY"] = True

UPLOAD_FOLDER = "uploads"
TEMPLATE_FOLDER = "templates_docs"

app.config['SQLALCHEMY_DATABASE_URI'] = os.environ.get("DATABASE_URL", "sqlite:///documents.db")
db = SQLAlchemy(app)

socketio = SocketIO(app, async_mode="threading")

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"
login_manager.login_message = None

# ---------------- USER MODEL ----------------
class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100))
    uid = db.Column(db.String(50))

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

# ---------------- DOCUMENT MODEL ----------------
class Document(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    filename = db.Column(db.String(200))
    reference = db.Column(db.String(100))
    status = db.Column(db.String(100))
    doc_type = db.Column(db.String(50))
    signer1 = db.Column(db.String(100))
    signer2 = db.Column(db.String(100))
    created_at = db.Column(db.DateTime, default=lambda: datetime.datetime.now())  # ✅ fixed time
    completed_at = db.Column(db.DateTime, nullable=True)

# ---------------- LOGIN ----------------
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        name = request.form.get('name')
        uid = request.form.get('uid')
        remember = True if request.form.get('remember') else False

        user = User.query.filter_by(uid=uid).first()
        if not user:
            user = User(name=name, uid=uid)
            db.session.add(user)
            db.session.commit()

        login_user(user, remember=remember)
        return redirect('/home')

    return render_template("login.html")

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect('/login')

# ---------------- HOME ----------------
@app.route('/')
def root():
    return redirect('/login')

@app.route('/home', methods=['GET', 'POST'])
@login_required
def home():

    if request.method == 'POST':

        action = request.form.get('action')
        reference = request.form.get('reference', '').strip()

        if not reference:
            return "Reference required"

        existing = Document.query.filter(
            Document.reference == reference,
            Document.status.contains("Awaiting")
        ).first()

        if existing:
            return f"Duplicate reference {reference} already awaiting signature."

        if action in ["qo1", "po1"]:

            template_file = "QO_Template.pdf" if action == "qo1" else "PO_Template.pdf"
            prefix = "QO" if action == "qo1" else "PO"

            template_path = os.path.join(TEMPLATE_FOLDER, template_file)
            if not os.path.exists(template_path):
                return f"Missing {template_file}"

            # Coordinates (unchanged from your provided version)
            if action == "qo1":
                x, y = 90, 505
            elif action == "po1":
                x, y = 110, 270

            signed_path = sign_pdf(
                template_path,
                current_user.name,
                current_user.uid,
                reference,
                x,
                y,
                prefix,
                "1st Sign"
            )

            doc = Document(
                filename=os.path.basename(signed_path),
                reference=reference,
                status=f"Awaiting {prefix} 2nd Sign",
                doc_type=prefix,
                signer1=current_user.name
            )

            db.session.add(doc)
            db.session.commit()

            return send_file(signed_path, as_attachment=True)

        elif action == "anfo1":

            pdf_file = request.files.get('pdf')
            if not pdf_file:
                return "Upload document"

            filepath = os.path.join(UPLOAD_FOLDER, pdf_file.filename)
            pdf_file.save(filepath)

            signed_path = sign_pdf(filepath, current_user.name, current_user.uid,
                                   reference, 90, 230, "ANFO", "1st Sign")

            doc = Document(
                filename=os.path.basename(signed_path),
                reference=reference,
                status="Awaiting ANFO 2nd Sign",
                doc_type="ANFO",
                signer1=current_user.name
            )

            db.session.add(doc)
            db.session.commit()

            return send_file(signed_path, as_attachment=True)

    return render_template("index.html", user=current_user)

# ---------------- DASHBOARD ----------------
@app.route('/dashboard')
@login_required
def dashboard():

    filter_type = request.args.get("filter", "open")
    search_query = request.args.get("search", "").strip()
    now = datetime.datetime.now()  # ✅ local time
    two_min_ago = now - datetime.timedelta(minutes=2)

    query = Document.query

    # ✅ SEARCH BY reference OR creator OR 2nd signer
    if search_query:
        query = query.filter(
            or_(
                Document.reference.contains(search_query),
                Document.signer1.contains(search_query),
                Document.signer2.contains(search_query)
            )
        )

    if filter_type == "open":
        query = query.filter(
            (Document.status.contains("Awaiting")) |
            (
                (Document.completed_at != None) &
                (Document.completed_at >= two_min_ago)
            )
        )

    elif filter_type == "completed":
        query = query.filter(Document.completed_at != None)

    elif filter_type == "qo":
        query = query.filter(Document.doc_type == "QO")

    elif filter_type == "po":
        query = query.filter(Document.doc_type == "PO")

    elif filter_type == "anfo":
        query = query.filter(Document.doc_type == "ANFO")

    documents = query.order_by(Document.created_at.desc()).all()

    pending_count = Document.query.filter(
        Document.status.contains("Awaiting")
    ).count()

    return render_template(
        "dashboard.html",
        documents=documents,
        pending_count=pending_count,
        filter_type=filter_type,
        search_query=search_query
    )

# ---------------- KPI ----------------
@app.route('/kpi')
@login_required
def kpi():

    period = request.args.get("period", "all")
    today = datetime.datetime.now()  # ✅ local time

    query = Document.query

    if period == "today":
        start = today.replace(hour=0, minute=0, second=0)
        query = query.filter(Document.created_at >= start)

    elif period == "week":
        start = today - datetime.timedelta(days=7)
        query = query.filter(Document.created_at >= start)

    elif period == "month":
        start = today - datetime.timedelta(days=30)
        query = query.filter(Document.created_at >= start)

    elif period == "quarter":
        start = today - datetime.timedelta(days=90)
        query = query.filter(Document.created_at >= start)

    elif period == "half":
        start = today - datetime.timedelta(days=180)
        query = query.filter(Document.created_at >= start)

    elif period == "year":
        start = today - datetime.timedelta(days=365)
        query = query.filter(Document.created_at >= start)

    docs = query.all()

    total_cases = len(docs)
    completed_cases = len([d for d in docs if d.completed_at])
    open_cases = total_cases - completed_cases

    total_minutes = 0
    for d in docs:
        if d.completed_at:
            total_minutes += (d.completed_at - d.created_at).total_seconds() / 60

    avg_time = round(total_minutes / completed_cases, 2) if completed_cases else 0

    qo_total = len([d for d in docs if d.doc_type == "QO"])
    po_total = len([d for d in docs if d.doc_type == "PO"])
    anfo_total = len([d for d in docs if d.doc_type == "ANFO"])

    users = User.query.all()
    user_stats = []

    for user in users:
        docs_1st = [d for d in docs if d.signer1 == user.name]
        docs_2nd = [d for d in docs if d.signer2 == user.name]

        stats = {
            "name": user.name,
            "qo_1st": len([d for d in docs_1st if d.doc_type == "QO"]),
            "qo_2nd": len([d for d in docs_2nd if d.doc_type == "QO"]),
            "po_1st": len([d for d in docs_1st if d.doc_type == "PO"]),
            "po_2nd": len([d for d in docs_2nd if d.doc_type == "PO"]),
            "anfo_1st": len([d for d in docs_1st if d.doc_type == "ANFO"]),
            "anfo_2nd": len([d for d in docs_2nd if d.doc_type == "ANFO"]),
            "total_1st": len(docs_1st),
            "total_2nd": len(docs_2nd),
        }

        stats["overall"] = stats["total_1st"] + stats["total_2nd"]
        user_stats.append(stats)

    return render_template("kpi.html",
                           total_cases=total_cases,
                           completed_cases=completed_cases,
                           open_cases=open_cases,
                           avg_time=avg_time,
                           qo_total=qo_total,
                           po_total=po_total,
                           anfo_total=anfo_total,
                           user_stats=user_stats,
                           period=period)

# ---------------- EXPORT ----------------
@app.route('/export')
@login_required
def export_excel():

    docs = Document.query.all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Detailed Cases"

    ws.append([
        "Reference","Type","Status","1st Signer","2nd Signer",
        "Created","Completed","Turnaround (Min)"
    ])

    for d in docs:
        turnaround = ""
        if d.completed_at:
            turnaround = round((d.completed_at - d.created_at).total_seconds() / 60, 2)

        ws.append([
            d.reference,
            d.doc_type,
            d.status,
            d.signer1,
            d.signer2,
            d.created_at.strftime('%d-%b-%Y %H:%M'),
            d.completed_at.strftime('%d-%b-%Y %H:%M') if d.completed_at else "",
            turnaround
        ])

    file_path = "enterprise_report.xlsx"
    wb.save(file_path)

    return send_file(file_path, as_attachment=True)

# ---------------- DELETE DOCUMENT ----------------
@app.route('/delete/<int:doc_id>')
@login_required
def delete_document(doc_id):

    document = Document.query.get(doc_id)

    if document:

        # Optional: remove file from uploads folder
        filepath = os.path.join(UPLOAD_FOLDER, document.filename)
        if os.path.exists(filepath):
            os.remove(filepath)

        db.session.delete(document)
        db.session.commit()

    return redirect('/dashboard')

# ---------------- SIGN SECOND ----------------
@app.route('/sign/<int:doc_id>')
@login_required
def sign_second(doc_id):

    document = Document.query.get(doc_id)

    if document and "Awaiting" in document.status:

        filepath = os.path.join(UPLOAD_FOLDER, document.filename)

        if document.doc_type == "QO":
            x, y = 399, 505
        elif document.doc_type == "PO":
            x, y = 410, 270
        else:
            x, y = 399, 230

        signed_path = sign_pdf(
            filepath,
            current_user.name,
            current_user.uid,
            document.reference,
            x,
            y,
            document.doc_type,
            "2nd Sign"
        )

        document.status = "Completed"
        document.signer2 = current_user.name
        document.filename = os.path.basename(signed_path)
        document.completed_at = datetime.datetime.now()  # ✅ fixed time

        db.session.commit()

        return send_file(signed_path, as_attachment=True)

    return redirect('/dashboard')

# ---------------- SIGN FUNCTION ----------------
def sign_pdf(filepath, name, uid, reference, x, y, prefix, stage):

    timestamp = datetime.datetime.now().strftime("%d-%b-%Y %H:%M")

    packet = io.BytesIO()
    can = canvas.Canvas(packet)
    can.setFont("Helvetica", 8)

    can.drawString(x, y, f"Signed by: {name}")
    can.drawString(x, y - 12, f"UID: {uid}")
    can.drawString(x, y - 24, f"Date: {timestamp}")
    can.drawString(x, y - 36, f"Ref: {reference}")

    can.save()
    packet.seek(0)

    overlay_pdf = PdfReader(packet)
    original_pdf = PdfReader(filepath)

    writer = PdfWriter()
    page = original_pdf.pages[0]
    page.merge_page(overlay_pdf.pages[0])
    writer.add_page(page)

    filename = f"{prefix} {reference} - {stage}.pdf"
    signed_path = os.path.join(UPLOAD_FOLDER, filename)

    with open(signed_path, "wb") as f:
        writer.write(f)

    return signed_path

# Ensure database tables exist
with app.app_context():
    db.create_all()


if __name__ == '__main__':
    if not os.path.exists(UPLOAD_FOLDER):
        os.makedirs(UPLOAD_FOLDER)
    if not os.path.exists(TEMPLATE_FOLDER):
        os.makedirs(TEMPLATE_FOLDER)
    with app.app_context():
        db.create_all()
    socketio.run(app)